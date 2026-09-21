"""Create non-destructive invoice output for CLI and future web callers."""

from __future__ import annotations

import shutil
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from .common import (
    EXCEL_HEADERS,
    EXPENSE_TYPES,
    InvoiceOrganizerError,
    RUNTIME_DIR,
    WORKBOOK_NAME,
    decimal_or_none,
    ensure_direct_child,
    input_snapshot,
    load_json,
    output_records,
    sha256_bytes,
    sha256_file,
    utc_now,
    validate_path_layout,
    write_json,
)


@dataclass(frozen=True)
class OrganizationResult:
    processed_count: int
    workbook_path: Path
    grand_total: Decimal | None
    plan_path: Path


def _money(value: Any) -> Decimal | None:
    return decimal_or_none(value)


def _record_row(record: dict[str, Any]) -> list[Any]:
    invoice_date = record.get("invoice_date")
    return [
        record["new_name"],
        record.get("invoice_number"),
        record.get("expense_type"),
        record.get("document_type"),
        _money(record.get("amount_excluding_tax")),
        record.get("tax_rate"),
        _money(record.get("tax_amount")),
        _money(record.get("fuel_surcharge")),
        _money(record.get("civil_aviation_development_fund")),
        _money(record.get("total_amount")),
        date.fromisoformat(invoice_date) if invoice_date else None,
        record.get("remark"),
    ]


def build_workbook(records: list[dict[str, Any]], target: Path, generated_at: str) -> Decimal | None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "发票整理清单"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    workbook.properties.creator = "invoice-organizer-skill"
    workbook.properties.created = datetime.fromisoformat(generated_at.replace("Z", "+00:00"))
    workbook.properties.modified = workbook.properties.created

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Aptos", size=10, color="222222")
    total_font = Font(name="Aptos", size=10, bold=True, color="222222")
    bottom_border = Border(bottom=Side(style="thin", color="B7C9D6"))

    sheet.append(list(EXCEL_HEADERS))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = bottom_border
    sheet.row_dimensions[1].height = 22

    valid_totals: list[Decimal] = []
    for record in records:
        sheet.append(_record_row(record))
        amount = _money(record.get("total_amount"))
        if amount is not None:
            valid_totals.append(amount)

    grand_total = sum(valid_totals, Decimal("0")) if valid_totals else None
    total_row = sheet.max_row + 1
    sheet.cell(total_row, 1, "合计")
    if grand_total is not None:
        sheet.cell(total_row, 10, grand_total)

    for row in sheet.iter_rows(min_row=2, max_row=total_row):
        for cell in row:
            cell.font = total_font if cell.row == total_row else body_font
            cell.alignment = Alignment(
                horizontal="right" if cell.column in {5, 7, 8, 9, 10} else "left",
                vertical="center",
            )
        for column in (5, 7, 8, 9, 10):
            row[column - 1].number_format = "0.00"
        row[10].number_format = "yyyy-mm-dd"
    for cell in sheet[total_row]:
        cell.border = Border(top=Side(style="thin", color="1F4E78"))

    widths = {
        "A": 46,
        "B": 24,
        "C": 14,
        "D": 30,
        "E": 14,
        "F": 10,
        "G": 12,
        "H": 14,
        "I": 14,
        "J": 14,
        "K": 14,
        "L": 24,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = f"A1:L{total_row}"
    workbook.save(target)
    return grand_total


def _read_output_payload(record: dict[str, Any], input_dir: Path) -> Path | bytes:
    if record["file_type"] == "invoice_pdf" or (
        record["file_type"] == "ride_report" and not record.get("archive_source")
    ):
        source = input_dir / record["original_name"]
        ensure_direct_child(source, input_dir)
        if not source.is_file():
            raise RuntimeError(f"input 源文件不存在: {record['original_name']}")
        if sha256_file(source) != record["source_sha256"]:
            raise RuntimeError(f"input 源文件在扫描后发生变化: {record['original_name']}")
        return source
    if record["file_type"] in {"archive_pdf", "ride_report"} and record.get("archive_source"):
        archive_path = input_dir / record["archive_source"]
        ensure_direct_child(archive_path, input_dir)
        if not archive_path.is_file():
            raise RuntimeError(f"input ZIP 不存在: {record['archive_source']}")
        if sha256_file(archive_path) != record["archive_sha256"]:
            raise RuntimeError(f"input ZIP 在扫描后发生变化: {record['archive_source']}")
        with zipfile.ZipFile(archive_path) as archive:
            data = archive.read(record["archive_member"])
        if sha256_bytes(data) != record["member_sha256"]:
            raise RuntimeError(f"ZIP PDF 成员在扫描后发生变化: {record['original_name']}")
        return data
    raise RuntimeError(f"计划包含不可输出的记录类型: {record.get('file_type')}")


def _preflight_output(
    plan: dict[str, Any],
    output_dir: Path,
    records: list[dict[str, Any]],
    payloads: dict[str, Path | bytes],
) -> tuple[bool, str | None]:
    expected_names = {record["new_name"] for record in records}
    if len(expected_names) != len(records):
        raise RuntimeError("collision: 计划中存在重复目标文件名。")

    if output_dir.exists():
        extra = {
            path.name
            for path in output_dir.iterdir()
            if path.name != ".gitkeep"
            and path.name not in expected_names
            and path.name != WORKBOOK_NAME
        }
        if extra:
            raise RuntimeError(
                "collision: output 中存在本计划之外的文件: " + ", ".join(sorted(extra))
            )

    for record in records:
        target = output_dir / record["new_name"]
        payload = payloads[record["new_name"]]
        expected_hash = (
            sha256_file(payload) if isinstance(payload, Path) else sha256_bytes(payload)
        )
        if expected_hash != record["output_sha256"]:
            raise RuntimeError(f"计划输出哈希不一致: {record['new_name']}")
        if target.exists() and sha256_file(target) != expected_hash:
            raise RuntimeError(f"collision: 同名输出文件内容不同: {target.name}")

    workbook = output_dir / WORKBOOK_NAME
    if workbook.exists():
        recorded_hash = plan.get("output_workbook_sha256")
        if not recorded_hash:
            raise RuntimeError(f"collision: 已存在来源不明的 {WORKBOOK_NAME}")
        current_hash = sha256_file(workbook)
        if current_hash != recorded_hash:
            raise RuntimeError(f"collision: {WORKBOOK_NAME} 与当前计划记录不一致")
        return False, current_hash
    return True, None


def organize_invoices(plan_path: Path, confirmation_token: str) -> OrganizationResult:
    """Execute an approved plan and return a structured result."""
    plan_path = plan_path.expanduser().resolve()
    if plan_path.parent != RUNTIME_DIR.resolve():
        raise InvoiceOrganizerError("plan 必须位于当前 Skill 的 .runtime 目录中。")
    if not plan_path.is_file():
        raise InvoiceOrganizerError(f"找不到 plan: {plan_path}")
    plan = load_json(plan_path)
    if plan.get("version") != 2:
        raise InvoiceOrganizerError("plan 版本不受支持，请重新运行 scan。")
    if confirmation_token != plan.get("approval_token"):
        raise InvoiceOrganizerError("审批令牌不匹配，请重新检查预览。")

    try:
        input_dir, output_dir = validate_path_layout(
            Path(plan["input_dir"]), Path(plan["output_dir"])
        )
    except ValueError as exc:
        raise InvoiceOrganizerError(str(exc)) from exc
    if str(RUNTIME_DIR.resolve()) != plan.get("runtime_dir"):
        raise InvoiceOrganizerError("plan 的 runtime 目录不属于当前 Skill。")
    if not input_dir.is_dir():
        raise InvoiceOrganizerError(f"input 目录不存在: {input_dir}")
    current_snapshot = input_snapshot(input_dir)
    if current_snapshot != plan.get("input_snapshot"):
        raise InvoiceOrganizerError("input 文件名、数量、大小或 SHA256 已变化，请重新运行 scan。")

    records = output_records(plan.get("records", []))
    if not records:
        raise InvoiceOrganizerError("计划中没有可输出的票据。")
    for record in records:
        expense_type = record.get("expense_type")
        if expense_type is not None and expense_type not in EXPENSE_TYPES:
            raise InvoiceOrganizerError(f"费用类型不受支持: {expense_type}")

    try:
        payloads = {
            record["new_name"]: _read_output_payload(record, input_dir)
            for record in records
        }
        create_workbook, existing_workbook_hash = _preflight_output(
            plan, output_dir, records, payloads
        )
    except RuntimeError as exc:
        raise InvoiceOrganizerError(str(exc)) from exc

    RUNTIME_DIR.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="invoice-workbook-", dir=RUNTIME_DIR) as temp_dir:
        workbook_temp = Path(temp_dir) / WORKBOOK_NAME
        grand_total: Decimal | None
        if create_workbook:
            grand_total = build_workbook(records, workbook_temp, plan["generated_at"])
            workbook_hash = sha256_file(workbook_temp)
        else:
            grand_total = (
                Decimal(plan["grand_total"]) if plan.get("grand_total") is not None else None
            )
            workbook_hash = existing_workbook_hash

        output_dir.mkdir(parents=True, exist_ok=True)
        for record in records:
            target = output_dir / record["new_name"]
            if target.exists():
                record["status"] = "skipped_identical"
                continue
            payload = payloads[record["new_name"]]
            if isinstance(payload, Path):
                shutil.copy2(payload, target)
                record["status"] = "copied"
            else:
                with target.open("xb") as stream:
                    stream.write(payload)
                record["status"] = "extracted"
            if sha256_file(target) != record["output_sha256"]:
                raise RuntimeError(f"输出校验失败: {target.name}")

        workbook_target = output_dir / WORKBOOK_NAME
        if create_workbook:
            if workbook_target.exists():
                raise RuntimeError(f"collision: 写入前出现同名文件: {WORKBOOK_NAME}")
            shutil.copy2(workbook_temp, workbook_target)
            if sha256_file(workbook_target) != workbook_hash:
                raise RuntimeError(f"输出校验失败: {WORKBOOK_NAME}")

    plan["approved"] = True
    plan["executed_at"] = utc_now()
    plan["output_workbook_sha256"] = workbook_hash
    plan["grand_total"] = str(grand_total) if grand_total is not None else None
    write_json(plan_path, plan)

    return OrganizationResult(
        processed_count=len(records),
        workbook_path=output_dir / WORKBOOK_NAME,
        grand_total=grand_total,
        plan_path=plan_path,
    )
