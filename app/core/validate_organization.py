"""Validate organized invoices for CLI and future web callers."""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from .common import (
    CATEGORY_ORDER,
    EXCEL_HEADERS,
    InvoiceOrganizerError,
    WORKBOOK_NAME,
    input_snapshot,
    load_json,
    output_records,
    read_pdf_record,
    sha256_file,
    validate_path_layout,
)


@dataclass(frozen=True)
class ValidationResult:
    input_count: int
    output_count: int
    grand_total: str | None
    errors: tuple[str, ...]

    @property
    def is_valid(self) -> bool:
        return not self.errors


def fail(errors: list[str], message: str) -> None:
    errors.append(message)


def as_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value))
    except InvalidOperation:
        return None


def as_iso_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def expected_cell(record: dict[str, Any], column: int) -> Any:
    keys = {
        2: "invoice_number",
        3: "expense_type",
        4: "document_type",
        5: "amount_excluding_tax",
        6: "tax_rate",
        7: "tax_amount",
        8: "fuel_surcharge",
        9: "civil_aviation_development_fund",
        10: "total_amount",
        11: "invoice_date",
        12: "remark",
    }
    return record.get(keys[column])


def validate_input(plan: dict[str, Any], input_dir: Path, errors: list[str]) -> None:
    expected = plan.get("input_snapshot", [])
    current = input_snapshot(input_dir)
    if current == expected:
        return
    expected_by_name = {item["name"]: item for item in expected}
    current_by_name = {item["name"]: item for item in current}
    for name in sorted(expected_by_name.keys() - current_by_name.keys()):
        fail(errors, f"input 文件缺失: {name}")
    for name in sorted(current_by_name.keys() - expected_by_name.keys()):
        fail(errors, f"input 出现新增文件: {name}")
    for name in sorted(expected_by_name.keys() & current_by_name.keys()):
        if current_by_name[name] != expected_by_name[name]:
            fail(errors, f"input 文件内容、大小或 SHA256 已变化: {name}")


def validate_output(
    plan: dict[str, Any],
    output_dir: Path,
    records: list[dict[str, Any]],
    errors: list[str],
) -> None:
    expected_names = {record["new_name"] for record in records}
    expected_files = expected_names | {WORKBOOK_NAME}
    if not output_dir.is_dir():
        fail(errors, f"output 目录不存在: {output_dir}")
        return
    actual_entries = {path.name for path in output_dir.iterdir() if path.name != ".gitkeep"}
    if actual_entries != expected_files:
        for name in sorted(expected_files - actual_entries):
            fail(errors, f"output 缺少文件: {name}")
        for name in sorted(actual_entries - expected_files):
            fail(errors, f"output 存在计划外文件: {name}")
    for path in output_dir.iterdir():
        if path.name == ".gitkeep":
            continue
        if path.is_dir():
            fail(errors, f"output 不允许包含目录: {path.name}")
        if path.suffix.lower() in {
            ".json",
            ".csv",
            ".md",
            ".zip",
            ".ofd",
            ".tmp",
            ".log",
        }:
            fail(errors, f"output 包含禁止的内部或源文件: {path.name}")
    for record in records:
        target = output_dir / record["new_name"]
        if target.is_file() and sha256_file(target) != record.get("output_sha256"):
            fail(errors, f"输出文件 SHA256 不匹配: {target.name}")


def validate_sequence_and_order(records: list[dict[str, Any]], errors: list[str]) -> None:
    invoices = [
        record
        for record in records
        if record.get("file_type") in {"invoice_pdf", "archive_pdf"}
    ]
    sequences: list[int] = []
    ordered_keys: list[tuple[int, str]] = []
    for record in invoices:
        match = re.match(r"^(\d+)、", record["new_name"])
        if not match:
            fail(errors, f"票据文件名缺少序号: {record['new_name']}")
            continue
        sequence = int(match.group(1))
        sequences.append(sequence)
        ordered_keys.append(
            (
                CATEGORY_ORDER.get(record.get("category", "其他"), 3),
                record.get("business_date") or "9999-99-99",
            )
        )
        if record.get("invoice_number") not in record["new_name"]:
            fail(errors, f"文件名缺少发票号码: {record['new_name']}")
        if f"{record.get('total_amount')}元" not in record["new_name"]:
            fail(errors, f"文件名缺少票面总金额: {record['new_name']}")
    if sorted(sequences) != list(range(1, len(invoices) + 1)):
        fail(errors, f"发票序号不连续或重复: {sorted(sequences)}")
    if ordered_keys != sorted(ordered_keys):
        fail(errors, "类别顺序或同类业务日期顺序不符合计划")

    by_name = {record["new_name"]: record for record in records}
    for report in (item for item in records if item.get("file_type") == "ride_report"):
        related = by_name.get(report.get("ride_relation"))
        if not related:
            continue
        report_seq = re.match(r"^(\d+)、", report["new_name"])
        invoice_seq = re.match(r"^(\d+)、", related["new_name"]) if related else None
        if not report_seq or not invoice_seq or report_seq.group(1) != invoice_seq.group(1):
            fail(errors, f"打车行程单与发票序号不一致: {report['new_name']}")


def validate_ride_source_fields(
    records: list[dict[str, Any]], input_dir: Path, errors: list[str]
) -> None:
    protected_fields = (
        "invoice_number",
        "amount_excluding_tax",
        "tax_rate",
        "tax_amount",
    )
    for record in records:
        if record.get("file_type") != "ride_report":
            continue
        try:
            if record.get("archive_source"):
                archive_path = input_dir / record["archive_source"]
                with zipfile.ZipFile(archive_path) as archive:
                    source_bytes = archive.read(record["archive_member"])
            else:
                source = input_dir / record["original_name"]
                source_bytes = source.read_bytes()
            extracted = read_pdf_record(record["original_name"], source_bytes)
        except Exception as exc:
            fail(errors, f"无法复核打车行程单自身字段: {record['original_name']}: {exc}")
            continue
        for field in protected_fields:
            if record.get(field) != extracted.get(field):
                fail(errors, f"打车行程单疑似继承发票字段 {field}: {record['new_name']}")


def validate_didi_source_fields(
    records: list[dict[str, Any]], input_dir: Path, errors: list[str]
) -> None:
    """Backward-compatible alias for the generalized ride-report validation."""
    validate_ride_source_fields(records, input_dir, errors)


def validate_workbook(
    plan: dict[str, Any],
    output_dir: Path,
    records: list[dict[str, Any]],
    errors: list[str],
) -> None:
    from openpyxl import load_workbook

    workbook_path = output_dir / WORKBOOK_NAME
    if not workbook_path.is_file():
        return
    recorded_hash = plan.get("output_workbook_sha256")
    if not recorded_hash:
        fail(errors, "plan 未记录 Excel SHA256")
    elif sha256_file(workbook_path) != recorded_hash:
        fail(errors, "Excel SHA256 与 plan 不一致")

    workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    if workbook.sheetnames != ["发票整理清单"]:
        fail(errors, f"Excel 工作表不符合要求: {workbook.sheetnames}")
        return
    sheet = workbook["发票整理清单"]
    headers = tuple(sheet.cell(1, column).value for column in range(1, 13))
    if headers != EXCEL_HEADERS:
        fail(errors, f"Excel 列结构不符合要求: {headers}")
    if sheet.freeze_panes != "A2":
        fail(errors, "Excel 未冻结第一行")
    if not sheet.auto_filter.ref:
        fail(errors, "Excel 未开启 AutoFilter")

    expected_rows = len(records) + 2
    if sheet.max_row != expected_rows:
        fail(errors, f"Excel 行数不匹配: 实际 {sheet.max_row}，应为 {expected_rows}")
        return
    if sheet.max_column != 12:
        fail(errors, f"Excel 列数不匹配: {sheet.max_column}")

    forbidden_blank_markers = {"None", "null", "N/A", "未知"}
    excel_totals: list[Decimal] = []
    for row_number, record in enumerate(records, start=2):
        filename = sheet.cell(row_number, 1).value
        if filename != record["new_name"]:
            fail(errors, f"Excel 文件名与 output 不一致: {filename!r}")
        for column in range(2, 13):
            actual = sheet.cell(row_number, column).value
            expected = expected_cell(record, column)
            if expected is None or expected == "":
                if actual is not None:
                    fail(
                        errors,
                        f"Excel 缺失字段未保持空白: {record['new_name']} / {EXCEL_HEADERS[column - 1]}",
                    )
                continue
            if isinstance(actual, str) and actual in forbidden_blank_markers:
                fail(errors, f"Excel 含禁止的空值占位符: {record['new_name']} / {actual}")
            if column in {5, 7, 8, 9, 10}:
                if as_decimal(actual) != as_decimal(expected):
                    fail(errors, f"Excel 金额字段不匹配: {record['new_name']} / {EXCEL_HEADERS[column - 1]}")
            elif column == 11:
                if as_iso_date(actual) != str(expected):
                    fail(errors, f"Excel 开票日期不匹配: {record['new_name']}")
            elif actual != expected:
                fail(errors, f"Excel 字段不匹配: {record['new_name']} / {EXCEL_HEADERS[column - 1]}")
        total = as_decimal(sheet.cell(row_number, 10).value)
        if total is not None:
            excel_totals.append(total)

    total_row = sheet.max_row
    if sheet.cell(total_row, 1).value != "合计":
        fail(errors, "Excel 最后一行第一列不是“合计”")
    for column in range(2, 13):
        if column == 10:
            continue
        if sheet.cell(total_row, column).value is not None:
            fail(errors, f"Excel 合计行第 {column} 列必须为空")
    expected_total = sum(excel_totals, Decimal("0")) if excel_totals else None
    actual_total = as_decimal(sheet.cell(total_row, 10).value)
    if actual_total != expected_total:
        fail(errors, f"Excel 票面总金额合计不正确: {actual_total} != {expected_total}")
    plan_total = as_decimal(plan.get("grand_total"))
    if plan_total != expected_total:
        fail(errors, f"plan 与 Excel 合计不一致: {plan_total} != {expected_total}")


def validate_organization(plan_path: Path) -> ValidationResult:
    """Validate an executed plan and return all validation errors."""
    plan_path = plan_path.expanduser().resolve()
    if not plan_path.is_file():
        raise InvoiceOrganizerError(f"找不到 plan: {plan_path}")
    plan = load_json(plan_path)
    if plan.get("version") != 2:
        raise InvoiceOrganizerError("plan 版本不受支持，请重新运行 scan。")
    try:
        input_dir, output_dir = validate_path_layout(
            Path(plan["input_dir"]), Path(plan["output_dir"])
        )
    except ValueError as exc:
        raise InvoiceOrganizerError(str(exc)) from exc

    errors: list[str] = []
    if not plan.get("approved"):
        fail(errors, "plan 尚未执行或未记录确认状态")
    validate_input(plan, input_dir, errors)
    records = output_records(plan.get("records", []))
    validate_output(plan, output_dir, records, errors)
    validate_sequence_and_order(records, errors)
    validate_ride_source_fields(records, input_dir, errors)
    validate_workbook(plan, output_dir, records, errors)

    return ValidationResult(
        input_count=len(plan["input_snapshot"]),
        output_count=len(records),
        grand_total=plan.get("grand_total"),
        errors=tuple(errors),
    )
