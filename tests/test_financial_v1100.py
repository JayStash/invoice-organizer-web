"""Expense totals are distinct from the amounts printed on attachments."""
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from app.core.common import EXCEL_HEADERS
from app.core.organize_invoices import build_workbook
from app.core.validate_organization import validate_expense_relations


def expense_records():
    invoice = {
        "original_name": "invoice.pdf", "new_name": "1、发票.pdf",
        "document_id": "input-a", "file_type": "invoice_pdf", "document_role": "invoice",
        "status": "planned", "total_amount": "30.00", "expense_id": "invoice-a",
        "counts_toward_total": True, "expense_amount": "30.00", "expense_basis": "invoice",
    }
    attachment = {
        "original_name": "trip.pdf", "new_name": "1、行程单.pdf",
        "document_id": "input-b", "file_type": "ride_report", "document_role": "itinerary",
        "status": "planned", "total_amount": "30.00", "expense_id": "invoice-a",
        "ride_relation": invoice["new_name"], "counts_toward_total": False,
        "expense_amount": None, "expense_basis": "attachment",
    }
    paper = {
        "original_name": "paper-trip.pdf", "new_name": "2、行程单（对应纸质发票）.pdf",
        "document_id": "input-c", "file_type": "ride_report", "document_role": "itinerary",
        "status": "planned", "total_amount": "12.50", "expense_id": "paper-c",
        "standalone_paper_invoice": True, "counts_toward_total": True,
        "expense_amount": "12.50", "expense_basis": "paper_itinerary",
    }
    receipt = {
        "original_name": "receipt.pdf", "new_name": "3、货运收款凭证.pdf",
        "document_id": "input-d", "file_type": "supporting_receipt", "document_role": "supporting_receipt",
        "status": "planned", "total_amount": "99.00", "expense_id": None,
        "counts_toward_total": False, "expense_amount": None, "expense_basis": "supporting_receipt",
    }
    return [invoice, attachment, paper, receipt]


class FinancialV1100Tests(unittest.TestCase):
    def test_twelve_columns_keep_document_values_but_total_counts_each_expense_once(self):
        records = expense_records()
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "list.xlsx"
            total = build_workbook(records, path, "2026-10-23T00:00:00Z")
            self.assertEqual(total, Decimal("42.50"))
            workbook = load_workbook(path, data_only=True)
            try:
                sheet = workbook.active
                self.assertEqual(sheet.max_column, 12)
                self.assertEqual(tuple(cell.value for cell in sheet[1]), EXCEL_HEADERS)
                self.assertEqual([sheet.cell(n, 10).value for n in range(2, 6)], [30, 30, 12.5, 99])
                self.assertEqual(sheet.cell(6, 10).value, 42.5)
                self.assertIn("独立费用", sheet.cell(6, 10).comment.text)
            finally:
                workbook.close()
        errors = []
        validate_expense_relations(records, errors)
        self.assertEqual(errors, [])

    def test_validator_rejects_charged_attachment_even_if_producer_flags_claim_it_is_valid(self):
        records = expense_records()
        records[1].update(counts_toward_total=True, expense_amount="30.00")
        errors = []
        validate_expense_relations(records, errors)
        self.assertTrue(any("票据角色" in message for message in errors), errors)

    def test_validator_rejects_missing_invoice_charge_and_duplicate_expense_identity(self):
        records = expense_records()
        records[0].update(counts_toward_total=False, expense_amount=None)
        duplicate = deepcopy(records[0])
        duplicate.update(document_id="other-input", original_name="other.pdf", new_name="4、其他.pdf",
                         counts_toward_total=True, expense_amount="30.00")
        records.append(duplicate)
        errors = []
        validate_expense_relations(records, errors)
        self.assertTrue(any("票据角色" in message for message in errors), errors)
        self.assertTrue(any("多次计入" in message for message in errors), errors)

    def test_reviewable_known_itinerary_preserves_own_amount_and_contributes_nothing(self):
        record = expense_records()[1]
        record.update(status="needs_review", ride_relation=None, new_name=None,
                      expense_id=None, expense_basis="unresolved")
        errors = []
        validate_expense_relations([record], errors)
        self.assertEqual(errors, [])
        self.assertEqual(record["total_amount"], "30.00")


if __name__ == "__main__":
    unittest.main()
