"""Privacy-safe regressions for Web batch isolation and cleanup failures."""

from __future__ import annotations

import importlib
import io
import stat
import tempfile
import unittest
import zipfile
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException, UploadFile
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app.main as web
from app import runtime
from app.core import InvoiceOrganizerError
from test_invoice_v131 import NUMBER, synthetic_aviation_pdf


class WebCleanupTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve()
        self.readonly_files: list[Path] = []
        self.addCleanup(self._restore_writable_files)
        self.patches = ExitStack()
        self.addCleanup(self.patches.close)

        # configure_web_runtime changes core module globals. Restore every one
        # after each test so neither other tests nor source data inherit them.
        core_globals = {
            "app.core.common": ("SKILL_ROOT", "DEFAULT_INPUT_DIR", "DEFAULT_OUTPUT_DIR",
                                "RUNTIME_DIR", "DEFAULT_PLAN_PATH", "DEFAULT_PREVIEW_PATH"),
            "app.core.scan_invoices": ("DEFAULT_INPUT_DIR", "DEFAULT_OUTPUT_DIR", "RUNTIME_DIR",
                                       "DEFAULT_PLAN_PATH", "DEFAULT_PREVIEW_PATH"),
            "app.core.organize_invoices": ("RUNTIME_DIR",),
        }
        for module_name, names in core_globals.items():
            module = importlib.import_module(module_name)
            self.patches.enter_context(patch.multiple(
                module, **{name: getattr(module, name) for name in names}
            ))
        self.paths = runtime.configure_web_runtime(self.root)
        self.patches.enter_context(patch.multiple(
            web,
            RUNTIME_PATHS=self.paths,
            PROJECT_ROOT=self.paths.root,
            INPUT_DIR=self.paths.input_dir,
            OUTPUT_DIR=self.paths.output_dir,
            RUNTIME_DIR=self.paths.runtime_dir,
            DEFAULT_PLAN_PATH=self.paths.plan_path,
            PREVIEW_PATH=self.paths.preview_path,
            ZIP_PATH=self.paths.runtime_dir / web.DOWNLOAD_NAME,
            _active_batch=None,
        ))
        self.client = TestClient(web.app)
        self.addCleanup(self.client.close)

    def _restore_writable_files(self) -> None:
        for path in self.readonly_files:
            if path.exists():
                path.chmod(stat.S_IREAD | stat.S_IWRITE)

    def _write(self, relative: str, contents: bytes, *, readonly: bool = False) -> Path:
        path = self.root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(contents)
        if readonly:
            path.chmod(stat.S_IREAD)
            self.readonly_files.append(path)
        return path

    def _historical_files(self) -> dict[Path, bytes]:
        preserved = {
            "input/source-invoice.pdf": b"original CLI input",
            "output/source-results.xlsx": b"original CLI results",
            ".runtime/v131/real-regression-test/historical-group-1-flow/input/history.pdf": b"historical sample",
            ".runtime/v131/release-report.md": b"published release evidence",
            ".runtime/.invoice-organization-plan.json": b"CLI plan",
            ".runtime/发票整理预览.md": b"CLI preview",
        }
        return {self._write(name, data, readonly=True): data for name, data in preserved.items()}

    def _assert_preserved(self, preserved: dict[Path, bytes]) -> None:
        for path, data in preserved.items():
            self.assertEqual(path.read_bytes(), data, str(path))

    def test_cleanup_only_removes_web_owned_leaf_files(self) -> None:
        preserved = self._historical_files()
        for relative in (
            "input/.gitkeep", "input/notes.txt", "input/nested/history.pdf",
            "input/retained.pdf/child.pdf", "output/notes.txt", "output/nested/history.pdf",
            ".runtime/release-report.md", ".runtime/unknown.zip", ".runtime/nested/history.pdf",
        ):
            path = self._write(f".runtime/web-batch/{relative}", b"must survive", readonly=True)
            preserved[path] = b"must survive"
        disposable = []
        for relative in ("input/invoice.PDF", "input/invoices.zip", "output/invoice.pdf",
                         "output/results.XLSX", ".runtime/.invoice-organization-plan.json",
                         ".runtime/发票整理预览.md", f".runtime/{web.DOWNLOAD_NAME}"):
            disposable.append(self._write(f".runtime/web-batch/{relative}", b"old web batch"))

        for _ in range(2):
            web._clear_managed_data()
            self._assert_preserved(preserved)
            self.assertTrue(all(not path.exists() for path in disposable))
            self.assertTrue((self.paths.input_dir / "retained.pdf").is_dir())

    def test_cleanup_refuses_an_unmanaged_directory(self) -> None:
        path = self._write("input/cli.pdf", b"CLI data")
        with self.assertRaises(web.BatchCleanupError):
            web._clear_directory(path.parent)
        self.assertEqual(path.read_bytes(), b"CLI data")

    def test_cleanup_refuses_linked_roots_directories_and_leaf_files(self) -> None:
        leaf = self.paths.input_dir / "linked.pdf"
        leaf.write_bytes(b"external target remains safe")
        plan = self.paths.plan_path
        plan.write_text("external plan remains safe", encoding="utf-8")
        candidates = (self.paths.root, self.paths.input_dir, self.paths.output_dir,
                      self.paths.runtime_dir, leaf, plan)
        for linked in candidates:
            with (
                self.subTest(linked=linked),
                patch.object(web, "is_link_or_junction", side_effect=lambda path: path == linked),
                self.assertRaises(web.BatchCleanupError),
            ):
                web._clear_managed_data()
            # A previous subcase may clear ordinary leaves before encountering
            # another protected directory. Recreate only test-owned fixtures.
            if linked == leaf:
                self.assertEqual(leaf.read_bytes(), b"external target remains safe")
            if linked == plan:
                self.assertEqual(plan.read_text(encoding="utf-8"), "external plan remains safe")
            leaf.write_bytes(b"external target remains safe")
            plan.write_text("external plan remains safe", encoding="utf-8")

    def test_initial_permission_failure_is_not_retried_and_upload_is_closed(self) -> None:
        upload = UploadFile(filename="invoice.pdf", file=io.BytesIO(b"synthetic"))
        denied = PermissionError(5, "Access is denied", str(self.paths.input_dir / "locked.pdf"))
        with (
            patch.object(web, "_clear_managed_data", side_effect=denied) as cleanup,
            patch.object(web, "scan_invoices") as scan,
            self.assertLogs(web.logger, level="ERROR"),
            self.assertRaises(HTTPException) as raised,
        ):
            web.scan_uploads([upload])
        self.assertEqual(raised.exception.status_code, 409)
        self.assertIs(raised.exception.__cause__, denied)
        cleanup.assert_called_once_with()
        scan.assert_not_called()
        self.assertTrue(upload.file.closed)

    def test_initial_unexpected_cleanup_failure_is_not_retried(self) -> None:
        upload = UploadFile(filename="invoice.pdf", file=io.BytesIO(b"synthetic"))
        original = RuntimeError("synthetic cleanup implementation error")
        with (
            patch.object(web, "_clear_managed_data", side_effect=original) as cleanup,
            patch.object(web, "scan_invoices") as scan,
            self.assertLogs(web.logger, level="ERROR"),
            self.assertRaises(HTTPException) as raised,
        ):
            web.scan_uploads([upload])
        self.assertEqual(raised.exception.status_code, 409)
        self.assertIs(raised.exception.__cause__, original)
        cleanup.assert_called_once_with()
        scan.assert_not_called()
        self.assertTrue(upload.file.closed)

    def test_cleanup_failure_preserves_identification_and_unexpected_errors(self) -> None:
        for original, status in ((InvoiceOrganizerError("synthetic invoice cannot be recognized"), 400),
                                 (ValueError("synthetic unexpected error"), 500)):
            upload = UploadFile(filename="invoice.pdf", file=io.BytesIO(b"synthetic"))
            with (
                self.subTest(status=status),
                patch.object(web, "_clear_managed_data", side_effect=[None, PermissionError("locked")]) as cleanup,
                patch.object(web, "scan_invoices", side_effect=original),
                self.assertLogs(web.logger, level="ERROR") as logged,
                self.assertRaises(HTTPException) as raised,
            ):
                web.scan_uploads([upload])
            self.assertEqual(raised.exception.status_code, status)
            self.assertIs(raised.exception.__cause__, original)
            self.assertEqual(cleanup.call_count, 2)
            self.assertTrue(upload.file.closed)
            self.assertTrue(any("保留上传或识别的原始异常" in message for message in logged.output))
            # The fake cleanup left a partial upload. A real next cleanup must
            # make the same filename available again without user intervention.
            web._clear_managed_data()

    def test_upload_copy_and_close_failures_preserve_the_copy_error(self) -> None:
        stream = io.BytesIO(b"synthetic")
        upload = UploadFile(filename="invoice.pdf", file=stream)
        original = OSError("synthetic upload copy failure")
        with (
            patch.object(web.shutil, "copyfileobj", side_effect=original),
            patch.object(stream, "close", side_effect=OSError("synthetic close failure")),
            patch.object(web, "scan_invoices") as scan,
            self.assertLogs(web.logger, level="ERROR"),
            self.assertRaises(HTTPException) as raised,
        ):
            web.scan_uploads([upload])
        stream.close()
        self.assertEqual(raised.exception.status_code, 500)
        self.assertIs(raised.exception.__cause__, original)
        scan.assert_not_called()
        self.assertEqual(list(self.paths.input_dir.iterdir()), [])

    def test_failed_request_then_two_real_pdf_batches_can_organize_and_download(self) -> None:
        preserved = self._historical_files()
        with (
            patch.object(web, "scan_invoices", side_effect=InvoiceOrganizerError("synthetic failure")),
            self.assertLogs(web.logger, level="ERROR"),
        ):
            response = self.client.post("/api/scan", files={"files": ("invoice.pdf", b"synthetic", "application/pdf")})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(list(self.paths.input_dir.iterdir()), [])
        previous_download = None
        pdf = synthetic_aviation_pdf()
        for index in range(2):
            with self.subTest(batch=index + 1):
                scan = self.client.post("/api/scan", files={"files": ("invoice.pdf", pdf, "application/pdf")})
                self.assertEqual(scan.status_code, 200, scan.text)
                payload = scan.json()
                self.assertEqual(payload["planned_count"], 1)
                self.assertEqual(payload["review_count"], 0)
                self.assertEqual(payload["records"][0]["total_amount"], "448.86")
                self.assertEqual(payload["records"][0]["invoice_number"], NUMBER)
                if previous_download:
                    self.assertEqual(self.client.get(previous_download).status_code, 404)
                organize = self.client.post("/api/organize", json={"batch_id": payload["batch_id"]})
                self.assertEqual(organize.status_code, 200, organize.text)
                self.assertEqual(organize.json()["processed_count"], 1)
                self.assertEqual(organize.json()["grand_total"], "448.86")
                previous_download = organize.json()["download_url"]
                downloaded = self.client.get(previous_download)
                self.assertEqual(downloaded.status_code, 200)
                with zipfile.ZipFile(io.BytesIO(downloaded.content)) as archive:
                    self.assertEqual(set(archive.namelist()), {
                        f"1、机票-{NUMBER}-448.86元.pdf", "发票整理清单.xlsx", web.FEEDBACK_FILENAME,
                    })
                    self.assertEqual(archive.read(f"1、机票-{NUMBER}-448.86元.pdf"), pdf)
                    workbook = load_workbook(io.BytesIO(archive.read("发票整理清单.xlsx")), data_only=True)
                    self.assertEqual(workbook.active.cell(2, 10).value, 448.86)
                    workbook.close()
                self._assert_preserved(preserved)


if __name__ == "__main__":
    unittest.main()
